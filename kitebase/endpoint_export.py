"""
kitebase.endpoint_export — export the rows of a view to a file.

Two endpoints, one gesture. `export` receives the situation of the view and
answers with the menu of formats; each entry carries the same situation back
with the format, and the client sends it to `export_file` through the download
route, which answers with the file.

The situation is what the person sees: the ticked rows, or the rows on screen,
as ids in the order of the grid, plus the visible columns with their titles.
The rows are re-read here with the view's own query and `pk in ids` added, so
an id the view would not show stays out; the order is the client's, because
order is presentation, not permission.

    {
      "query":   {...},                       # the view's query, as buildQuery makes it
      "ids":     [3, 1, 2],                   # in the order of the grid
      "columns": [{"field": "title", "title": "Title"}, ...],
      "format":  "xlsx" | "csv" | "tsv" | "json"   # export_file only
    }

The rows are read in blocks of ids, written and let go: what stays in memory
is the file and one block. The number of rows has a ceiling, the `dataview`
key `export_max_rows` of config.yaml: an export is for working on what one
sees in a spreadsheet, and beyond that it is a different tool.
"""
import csv
import datetime
import decimal
import io
import json
from typing import Any, Dict, Iterator, List, Optional, Sequence

import kitebase
from kitebase.apptime import today
from kitebase.endpoint_db import pk_field
from kitebase.endpoints import FileResult, endpoint
from kitebase.i18n import _, _f
from kitebase.querybuilder import DynamicQueryBuilder, JSONEncoder

DEFAULT_MAX_ROWS = 50_000

# Ids per query: small enough for any database's limit on bound parameters,
# large enough that the number of queries does not matter.
BLOCK = 500

FORMATS = {
    'xlsx': {'label': 'Excel (.xlsx)', 'ext': 'xlsx',
             'mime': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'},
    'csv':  {'label': 'CSV (.csv)', 'ext': 'csv', 'mime': 'text/csv'},
    'tsv':  {'label': 'Text, tab separated (.txt)', 'ext': 'txt',
             'mime': 'text/tab-separated-values'},
    'json': {'label': 'JSON (.json)', 'ext': 'json', 'mime': 'application/json'},
}


def max_rows() -> int:
    app = kitebase.utils.get_app()
    return int(app.pm.config.get('dataview', {}).get('export_max_rows', DEFAULT_MAX_ROWS))


def _situation(data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Check what the client sent, and refuse what cannot be exported.

    Raises ValueError with a message for the person: the endpoints turn it
    into a refusal.
    """
    query = data.get('query')
    if not isinstance(query, dict) or not query.get('table'):
        raise ValueError(_('Query not defined'))
    ids = data.get('ids')
    if not isinstance(ids, list) or not ids:
        raise ValueError(_('No rows to export'))
    columns = data.get('columns')
    if not isinstance(columns, list) or not columns or not all(
            isinstance(c, dict) and isinstance(c.get('field'), str) for c in columns):
        raise ValueError(_('No columns to export'))
    limit = max_rows()
    if len(ids) > limit:
        raise ValueError(_f('{n} rows are more than the {limit} an export can hold: '
                            'narrow the selection', n=len(ids), limit=limit))
    return {'query': query, 'ids': ids, 'columns': columns}


@endpoint('export')
def export_menu(data: Dict[str, Any]) -> Dict[str, Any]:
    """The menu of formats, each entry carrying the situation back."""
    try:
        situation = _situation(data)
    except ValueError as e:
        return {'status': 'error', 'message': str(e), 'code': 400}

    n = len(situation['ids'])
    return {
        'status': 'success',
        'data': {
            'action': 'choose',
            'title': _('Export'),
            'message': _f('{n} rows', n=n) if n != 1 else _('1 row'),
            'options': [
                {'label': _(spec['label']),
                 'then': {'action': 'download', 'op': 'export_file',
                          'params': {**situation, 'format': fmt}}}
                for fmt, spec in FORMATS.items()
            ],
        },
    }


@endpoint('export_file')
def export_file(data: Dict[str, Any]) -> Any:
    """The file, through the download route."""
    fmt = data.get('format')
    if fmt not in FORMATS:
        return {'status': 'error', 'code': 400,
                'message': _f("Unknown export format: '{format}'", format=fmt)}
    try:
        situation = _situation(data)
    except ValueError as e:
        return {'status': 'error', 'message': str(e), 'code': 400}

    query, ids, columns = situation['query'], situation['ids'], situation['columns']
    table = query['table']
    app = kitebase.utils.get_app()
    pk = pk_field(app.tables.get(table))

    # The key goes first, always: it is what an import names a row by.
    fields = [c['field'] for c in columns if c['field'] != pk]
    headers = [pk] + [str(c.get('title') or c['field']) for c in columns if c['field'] != pk]
    keys = [pk] + [_key(f) for f in fields]

    with app.get_session() as session:
        rows = _rows(DynamicQueryBuilder(session, app.models), query, pk, [pk] + fields, ids)
        content = WRITERS[fmt](table, headers, keys, rows)

    spec = FORMATS[fmt]
    return FileResult(content, f'{table.lower()}_{today():%Y%m%d}.{spec["ext"]}', spec['mime'])


def _key(field: str) -> str:
    """The name a row carries a field under: what the grid calls it too
    (`extractFieldKey` of the client) - the alias, or the column's own name."""
    lower = field.lower()
    if ' as ' in lower:
        return field[lower.rindex(' as ') + 4:].strip()
    return field.rsplit('.', 1)[-1]


def _rows(builder: DynamicQueryBuilder, query: Dict[str, Any], pk: str,
          select: List[str], ids: Sequence[Any]) -> Iterator[Sequence[Any]]:
    """
    The rows of `ids` that the view's query admits, in the order of `ids`.

    One query per block. The view's conditions stay a group of their own and
    the block is a sibling group, so they are ANDed whatever the view's group
    opens with (an ['op', 'or'] would otherwise take the block in); ordering
    and pagination are dropped, because the order is the ids' and the set is
    the ids'.
    """
    base = {k: v for k, v in query.items()
            if k not in ('select', 'filters', 'order_by', 'limit', 'offset')}
    inner = (query.get('filters') or {}).get('conditions')
    view = inner if isinstance(inner, list) else ([inner] if inner else [])

    for start in range(0, len(ids), BLOCK):
        block = ids[start:start + BLOCK]
        only = [{pk: ['in', block]}]
        q = {**base, 'select': select,
             'filters': {'conditions': [view, only] if view else only}}
        by_id = {row[0]: row for row in builder.execute_query(q)}
        for i in block:
            row = by_id.get(i)
            if row is not None:
                yield row


# ── Writers ───────────────────────────────────────────────────────────────────
# (table, headers, keys, rows) -> bytes. The values arrive as the database gives
# them. Headers are the grid's titles, for a person; keys are the field names,
# for a program.

def _cell(value: Any) -> Any:
    """A value openpyxl can write as a typed cell."""
    if isinstance(value, datetime.datetime) and value.tzinfo is not None:
        # Excel has no timezones: the time as the database holds it.
        return value.replace(tzinfo=None)
    if isinstance(value, (str, int, float, bool, decimal.Decimal,
                          datetime.date, datetime.time)) or value is None:
        return value
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _text(value: Any) -> str:
    """A value as text, the form the endpoints accept back on write."""
    if value is None:
        return ''
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _delimited(delimiter: str):
    def write(table: str, headers: List[str], keys: List[str],
              rows: Iterator[Sequence[Any]]) -> bytes:
        buf = io.StringIO()
        out = csv.writer(buf, delimiter=delimiter, lineterminator='\r\n')
        out.writerow(headers)
        for row in rows:
            out.writerow([_text(v) for v in row])
        # The BOM is what makes Excel read the file as UTF-8.
        return buf.getvalue().encode('utf-8-sig')
    return write


def _xlsx(table: str, headers: List[str], keys: List[str],
          rows: Iterator[Sequence[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=table[:31])
    ws.freeze_panes = 'A2'
    for i, title in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, len(title) + 2)

    bold = Font(bold=True)
    header = []
    for title in headers:
        cell = WriteOnlyCell(ws, value=title)
        cell.font = bold
        header.append(cell)
    ws.append(header)
    for row in rows:
        ws.append([_cell(v) for v in row])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _json(table: str, headers: List[str], keys: List[str],
          rows: Iterator[Sequence[Any]]) -> bytes:
    """An array of objects keyed by field name, encoded as the API encodes.

    Written a row at a time, like the others: no list of dicts in memory.
    """
    buf = io.StringIO()
    buf.write('[')
    for n, row in enumerate(rows):
        buf.write(',\n' if n else '\n')
        buf.write(json.dumps(dict(zip(keys, row)), cls=JSONEncoder, ensure_ascii=False))
    buf.write('\n]\n')
    return buf.getvalue().encode('utf-8')


WRITERS = {'xlsx': _xlsx, 'csv': _delimited(','), 'tsv': _delimited('\t'), 'json': _json}
