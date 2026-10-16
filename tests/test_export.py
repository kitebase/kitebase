"""Exporting the rows of a view, on a database.

The client sends ids in the order of the grid and the view's own query; the
file has those rows, in that order, and only those the query admits. The
values keep their types: a spreadsheet that receives dates as text is one
where nothing can be sorted or summed.
"""
import datetime
import decimal
import importlib.util
import io

import pytest
import yaml
from openpyxl import load_workbook

import kitebase.endpoint_export as export
import kitebase.utils
from kitebase.db import DB, Base
from kitebase.endpoints import FileResult
from kitebase.plugins import PluginsManager
from kitebase.source import Generator

TABLES = {
    'Author': {
        'columns': [
            {'name': 'id', 'type': 'Integer', 'primary_key': True, 'autoincrement': True},
            {'name': 'name', 'type': 'String', 'length': 60},
        ],
    },
    'Book': {
        'columns': [
            {'name': 'id', 'type': 'Integer', 'primary_key': True, 'autoincrement': True},
            {'name': 'title', 'type': 'String', 'length': 80},
            {'name': 'price', 'type': 'Numeric', 'precision': 10, 'scale': 2},
            {'name': 'published', 'type': 'Date'},
            {'name': 'author_id', 'type': 'Integer', 'foreign_key': {'target': 'Author.id'}},
        ],
    },
}

BOOKS = [
    # title, price, published, author
    ('Kindred', '12.50', datetime.date(1979, 6, 1), 'Butler'),
    ('Dune', '9.90', datetime.date(1965, 8, 1), 'Herbert'),
    ('Solaris', '11.00', datetime.date(1961, 1, 1), 'Lem'),
    ('Ubik', '8.00', datetime.date(1969, 5, 1), 'Dick'),
]


@pytest.fixture(autouse=True)
def fresh_registry():
    yield
    Base.registry.dispose()
    Base.metadata.clear()


def build_app(tmp_path, monkeypatch, app_config=None):
    plugin = tmp_path / 'plugins' / 'shelf'
    plugin.mkdir(parents=True)
    (plugin / 'config.yaml').write_text(yaml.safe_dump({'name': 'shelf', 'version': '0.0.1'}))
    (plugin / 'model.yaml').write_text(yaml.safe_dump({'tables': TABLES}))

    cfg = tmp_path / 'config.yaml'
    cfg.write_text(yaml.safe_dump({'name': 'shelf', 'plugins': ['plugins'], **(app_config or {})}))
    monkeypatch.chdir(tmp_path)

    pm = PluginsManager()
    pm.load_config(str(cfg))
    kitebase.utils.register_standard_handlers(pm)
    pm.load_plugins()

    instance = DB()
    instance.calc_db(pm)

    source = tmp_path / 'model.py'
    Generator(instance).generate(filename=str(source))
    spec = importlib.util.spec_from_file_location('shelf_model', source)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    instance.initialize_db(f'sqlite:///{tmp_path / "shelf.sqlite"}', module, check_schema=False)
    monkeypatch.setattr(kitebase.utils, 'get_app', lambda: instance)

    Author, Book = instance.find_model_class('Author'), instance.find_model_class('Book')
    with instance.get_session() as session:
        for title, price, published, name in BOOKS:
            author = Author(name=name)
            session.add(author)
            session.flush()
            session.add(Book(title=title, price=decimal.Decimal(price),
                             published=published, author_id=author.id))
        session.commit()
    return instance


@pytest.fixture
def app(tmp_path, monkeypatch):
    return build_app(tmp_path, monkeypatch)


COLUMNS = [
    {'field': 'title', 'title': 'Title'},
    {'field': 'price', 'title': 'Price'},
    {'field': 'published', 'title': 'Published'},
]


def situation(ids, fmt='xlsx', columns=COLUMNS, query=None):
    return {'query': query or {'table': 'Book'}, 'ids': ids, 'columns': columns,
            'format': fmt}


def sheet(result):
    assert isinstance(result, FileResult), result
    ws = load_workbook(io.BytesIO(result.content)).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


# ── What goes in the file ─────────────────────────────────────────────────────

def test_the_rows_come_in_the_order_of_the_grid(app):
    rows = sheet(export.export_file(situation([3, 1, 4])))

    assert [r[1] for r in rows[1:]] == ['Solaris', 'Kindred', 'Ubik']


def test_the_key_comes_first_and_the_titles_are_the_grids(app):
    rows = sheet(export.export_file(situation([1])))

    assert rows[0] == ['id', 'Title', 'Price', 'Published']
    assert rows[1][0] == 1


def test_the_key_is_not_written_twice(app):
    columns = [{'field': 'id', 'title': 'Code'}] + COLUMNS
    rows = sheet(export.export_file(situation([1], columns=columns)))

    assert rows[0] == ['id', 'Title', 'Price', 'Published']


def test_the_values_keep_their_types(app):
    row = sheet(export.export_file(situation([1])))[1]

    assert row[2] == 12.5
    # openpyxl reads every date back as a datetime: a date cell, not text.
    assert row[3] == datetime.datetime(1979, 6, 1)


def test_a_joined_field_is_exported(app):
    query = {'table': 'Book', 'joins': [{'Author': 'Book.author_id = Author.id'}]}
    columns = COLUMNS[:1] + [{'field': 'Author.name', 'title': 'Author'}]
    rows = sheet(export.export_file(situation([2, 1], columns=columns, query=query)))

    assert rows[1:] == [[2, 'Dune', 'Herbert'], [1, 'Kindred', 'Butler']]


def test_an_id_the_view_does_not_admit_stays_out(app):
    """The ids are the client's; whether a row belongs to the view is not."""
    query = {'table': 'Book', 'filters': {'conditions': [{'price': ['>', 9]}]}}
    rows = sheet(export.export_file(situation([4, 1, 2], query=query)))

    assert [r[0] for r in rows[1:]] == [1, 2]


def test_a_view_that_opens_with_or_does_not_take_the_ids_in(app):
    """The ids are a group of their own: inside the view's OR they would
    widen it to every row the OR admits."""
    query = {'table': 'Book', 'filters': {'conditions': [
        'op', 'or', {'title': ['=', 'Dune']}, {'title': ['=', 'Ubik']}]}}
    rows = sheet(export.export_file(situation([1, 2], query=query)))

    assert [r[0] for r in rows[1:]] == [2]


def test_an_id_that_no_longer_exists_is_skipped(app):
    rows = sheet(export.export_file(situation([1, 99, 2])))

    assert [r[0] for r in rows[1:]] == [1, 2]


def test_the_order_holds_across_blocks(app, monkeypatch):
    monkeypatch.setattr(export, 'BLOCK', 2)
    rows = sheet(export.export_file(situation([4, 2, 3, 1])))

    assert [r[0] for r in rows[1:]] == [4, 2, 3, 1]


def test_the_views_order_and_paging_do_not_apply(app):
    query = {'table': 'Book', 'order_by': ['title'], 'limit': 1, 'offset': 2}
    rows = sheet(export.export_file(situation([3, 1], query=query)))

    assert [r[0] for r in rows[1:]] == [3, 1]


def test_csv_opens_in_excel_as_utf8(app):
    result = export.export_file(situation([1, 2], fmt='csv'))

    assert result.mime == 'text/csv'
    assert result.content.startswith(b'\xef\xbb\xbf')
    text = result.content.decode('utf-8-sig')
    assert text.splitlines() == ['id,Title,Price,Published',
                                 '1,Kindred,12.50,1979-06-01',
                                 '2,Dune,9.90,1965-08-01']


def test_tab_separated_text(app):
    result = export.export_file(situation([1], fmt='tsv'))

    assert result.filename.endswith('.txt')
    assert result.content.decode('utf-8-sig').splitlines()[1] == '1\tKindred\t12.50\t1979-06-01'


def test_json_is_keyed_by_field_name(app):
    """For a program: the field names, which do not change with the language,
    and the values as the API sends them."""
    import json
    result = export.export_file(situation([2, 1], fmt='json'))

    assert result.mime == 'application/json'
    assert json.loads(result.content) == [
        {'id': 2, 'title': 'Dune', 'price': 9.9, 'published': '1965-08-01'},
        {'id': 1, 'title': 'Kindred', 'price': 12.5, 'published': '1979-06-01'},
    ]


def test_json_names_a_joined_field_as_the_grid_does(app):
    import json
    query = {'table': 'Book', 'joins': [{'Author': 'Book.author_id = Author.id'}]}
    columns = [{'field': 'Author.name as author', 'title': 'Author'}]
    result = export.export_file(situation([2], fmt='json', columns=columns, query=query))

    assert json.loads(result.content) == [{'id': 2, 'author': 'Herbert'}]


def test_json_of_nothing_is_an_empty_array(app):
    import json
    result = export.export_file(situation([99], fmt='json'))

    assert json.loads(result.content) == []


def test_the_file_is_named_after_the_table_and_the_day(app):
    result = export.export_file(situation([1]))

    assert result.filename == f'book_{datetime.date.today():%Y%m%d}.xlsx'
    assert result.disposition == 'attachment'


# ── Refusals ──────────────────────────────────────────────────────────────────

def test_an_unknown_format_is_refused(app):
    result = export.export_file(situation([1], fmt='ods'))

    assert result['status'] == 'error'
    assert result['code'] == 400


@pytest.mark.parametrize('broken', [
    {'ids': []},
    {'columns': []},
    {'query': {}},
])
def test_an_incomplete_situation_is_refused(app, broken):
    result = export.export_file({**situation([1]), **broken})

    assert result['status'] == 'error'


def test_the_ceiling_is_the_installations(tmp_path, monkeypatch):
    build_app(tmp_path, monkeypatch, {'dataview': {'export_max_rows': 2}})

    result = export.export_file(situation([1, 2, 3]))

    assert result['status'] == 'error'
    assert '3' in result['message'] and '2' in result['message']


# ── The menu ──────────────────────────────────────────────────────────────────

def test_the_menu_carries_the_situation_back_with_the_format(app):
    sent = situation([3, 1])
    del sent['format']
    data = export.export_menu(sent)['data']

    assert data['action'] == 'choose'
    formats = [o['then']['params']['format'] for o in data['options']]
    assert formats == ['xlsx', 'csv', 'tsv', 'json']
    for option in data['options']:
        assert option['then']['action'] == 'download'
        assert option['then']['op'] == 'export_file'
        assert option['then']['params']['ids'] == [3, 1]


def test_the_menu_refuses_what_the_file_would(tmp_path, monkeypatch):
    build_app(tmp_path, monkeypatch, {'dataview': {'export_max_rows': 1}})

    assert export.export_menu(situation([1, 2]))['status'] == 'error'
